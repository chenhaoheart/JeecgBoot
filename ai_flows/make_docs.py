# -*- coding: utf-8 -*-
"""报告文档生成测试：章节 Markdown → 公文风格 HTML + Word(.docx)
模拟流程输出（Markdown 章节内容）→ 文档转换管线验证
"""
import json, re, os, datetime

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "output")
os.makedirs(OUT, exist_ok=True)
data = json.load(open(os.path.join(BASE, "_data_ch1.json"), encoding="utf8"))
ws, types, counties = data["ws"], data["types"], data["counties"]

def fmt(v):
    if v is None: return "-"
    s = str(v)
    try: return str(float(s)).rstrip("0").rstrip(".") if "." in s else s
    except Exception: return s

# ---- 组装章节 Markdown（与流程 LLM 输出同构）----
total_pop = sum(int(float(t["人口"] or 0)) for t in types)
total_obj = sum(int(t["数量"]) for t in types)
eles = sorted(float(w["沟口高程"]) for w in ws if w["沟口高程"])
md = []
md.append("# 一、流域概况\n")
md.append(f"本次山洪灾害风险隐患调查共涉及小流域 {len(ws)} 条，防治对象 {total_obj} 处，"
          f"涉及人口约 {total_pop:,} 人。防治对象以村落为主，分布集中于文县、武都区、徽县等县区，"
          "大部分防治对象沿沟道两岸分布，受山洪威胁较大。\n")
md.append(f"调查范围内小流域沟口高程介于 {fmt(eles[0])}~{fmt(eles[-1])} m 之间，沟道纵坡较大，"
          "洪水汇流时间短、涨幅陡，一旦发生强降雨，极易对沿岸防治对象形成威胁。\n")
md.append("调查范围内小流域基本信息见表 1-1。\n")
md.append("**表 1-1 小流域基本信息表**\n")
md.append("| 序号 | 流域名称 | 流域编码 | 所在沟道 | 所在位置 | 沟口高程(m) |")
md.append("| --- | --- | --- | --- | --- | --- |")
for i, w in enumerate(ws, 1):
    md.append(f"| {i} | {w['流域名称']} | {w['流域编码']} | {w['所在沟道'] or '-'} | {w['所在位置'] or '-'} | {fmt(w['沟口高程'])} |")
md.append("")
md.append(f"防治对象类型统计结果表明：村落 {types[0]['数量']} 处（人口 {int(float(types[0]['人口'])):,} 人），"
          f"占比最高；企事业单位共 {sum(t['数量'] for t in types if t['类型'] in ('企业','事业','事业单位'))} 处。"
          "各类型数量与人口分布见表 1-2。\n")
md.append("**表 1-2 防治对象类型统计表**\n")
md.append("| 类型 | 数量（处） | 人口（人） |")
md.append("| --- | --- | --- |")
for t in types:
    md.append(f"| {t['类型']} | {t['数量']} | {int(float(t['人口'] or 0)):,} |")
md.append("")
md.append("按县区统计的防治对象分布见表 1-3，其中武都区人口最为集中，防治任务较重。\n")
md.append("**表 1-3 县区防治对象分布表**\n")
md.append("| 县区 | 防治对象数（处） | 人口（人） |")
md.append("| --- | --- | --- |")
for c in counties:
    md.append(f"| {c['县区']} | {c['防治对象数']} | {int(float(c['人口'] or 0)):,} |")
md.append("")
CHAPTER_MD = "\n".join(md)
open(os.path.join(OUT, "chapter1.md"), "w", encoding="utf-8").write(CHAPTER_MD)

# ---- Markdown → 公文风格 HTML ----
def md_table_to_html(rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    h = "<table>\n<thead><tr>" + "".join(f"<th>{c}</th>" for c in header) + "</tr></thead>\n<tbody>"
    for r in body:
        h += "<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>"
    h += "</tbody></table>"
    return h

lines = CHAPTER_MD.split("\n")
html_body = []; i = 0; para = []
def flush_para():
    global para
    if para:
        html_body.append("<p>" + "".join(para) + "</p>"); para = []
while i < len(lines):
    ln = lines[i]
    if ln.startswith("| "):
        flush_para()
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(lines[i]); i += 1
        html_body.append(md_table_to_html(rows)); continue
    if ln.startswith("# "):
        flush_para(); html_body.append(f"<h1>{ln[2:]}</h1>")
    elif ln.startswith("**") and ln.endswith("**"):
        flush_para(); html_body.append(f"<p class='tbl-title'>{ln[2:-2]}</p>")
    elif ln.strip():
        para.append(ln.strip())
    else:
        flush_para()
    i += 1
flush_para()

HTML = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>山洪灾害风险隐患调查评价报告</title>
<style>
  body {{ font-family: "FangSong", "仿宋", "FangSong_GB2312", serif; font-size: 16pt; line-height: 1.8;
         color: #000; background: #f5f5f5; margin: 0; }}
  .paper {{ width: 210mm; min-height: 297mm; margin: 20px auto; background: #fff; padding: 25mm 22mm;
            box-shadow: 0 2px 12px rgba(0,0,0,.15); box-sizing: border-box; }}
  h1 {{ font-family: "SimHei", "黑体"; font-size: 18pt; text-align: left; margin: 24pt 0 12pt; }}
  .doc-title {{ font-family: "STZhongsong", "方正小标宋", "SimSun"; font-size: 22pt; text-align: center;
                margin: 0 0 8pt; }}
  .doc-sub {{ text-align: center; font-size: 12pt; color: #444; margin-bottom: 20pt; }}
  p {{ text-indent: 2em; margin: 0 0 10pt; }}
  .tbl-title {{ text-indent: 0; text-align: center; font-family: "SimHei", "黑体"; font-size: 12pt;
                margin: 14pt 0 6pt; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 10.5pt; font-family: "SimSun","宋体";
           margin: 0 0 14pt; }}
  th, td {{ border: 1px solid #000; padding: 4pt 6pt; text-align: center; }}
  th {{ background: #f0f0f0; font-family: "SimHei", "黑体"; font-weight: normal; }}
  @media print {{ body {{ background: #fff; }} .paper {{ margin: 0; box-shadow: none; }} }}
</style>
</head>
<body>
<div class="paper">
  <div class="doc-title">山洪灾害风险隐患调查评价报告</div>
  <div class="doc-sub">（AI 流程自动生成测试稿 · {datetime.date.today():%Y年%m月%d日}）</div>
  {"".join(html_body)}
</div>
</body>
</html>"""
html_path = os.path.join(OUT, "山洪灾害风险隐患调查评价报告.html")
open(html_path, "w", encoding="utf-8").write(HTML)
print("HTML:", html_path)

# ---- Markdown → Word (.docx) ----
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"; style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋_GB2312")
for s in doc.sections:
    s.top_margin, s.bottom_margin = Cm(2.54), Cm(2.54)
    s.left_margin, s.right_margin = Cm(2.8), Cm(2.6)

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("山洪灾害风险隐患调查评价报告"); r.font.size = Pt(22); r.font.name = "Times New Roman"
r.element.rPr.rFonts.set(qn("w:eastAsia"), "方正小标宋简体")
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run(f"（AI 流程自动生成测试稿 · {datetime.date.today():%Y年%m月%d日}）")
r.font.size = Pt(10.5)
if r.element.rPr is None:
    r.element.get_or_add_rPr()
if r.element.rPr.rFonts is None:
    r.element.rPr.get_or_add_rFonts()
r.element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")

def add_para(text, bold=False, hei=False, indent=True, size=12):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(28)
    run = p.add_run(text)
    run.font.size = Pt(size); run.font.bold = bold
    run.font.name = "Times New Roman"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体" if hei else "仿宋_GB2312")
    return p

lines = CHAPTER_MD.split("\n"); i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("| "):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(lines[i]); i += 1
        header = [c.strip().strip("*") for c in rows[0].strip("|").split("|")]
        body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
        tbl = doc.add_table(rows=1 + len(body), cols=len(header))
        tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            cell = tbl.rows[0].cells[j].paragraphs[0]; cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cell.add_run(h); run.font.bold = False; run.font.size = Pt(10.5)
            run.font.name = "Times New Roman"; run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        for ri, row in enumerate(body, 1):
            for j, v in enumerate(row):
                cell = tbl.rows[ri].cells[j].paragraphs[0]; cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cell.add_run(v); run.font.size = Pt(10.5)
                run.font.name = "Times New Roman"; run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        doc.add_paragraph(); continue
    if ln.startswith("# "):
        p = add_para(ln[2:], hei=True, size=16, indent=False)
    elif ln.startswith("**") and ln.endswith("**"):
        add_para(ln[2:-2], hei=True, size=12, indent=False)
    elif ln.strip():
        add_para(ln.strip())
    i += 1

docx_path = os.path.join(OUT, "山洪灾害风险隐患调查评价报告.docx")
doc.save(docx_path)
print("DOCX:", docx_path)
