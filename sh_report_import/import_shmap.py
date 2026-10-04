# -*- coding: utf-8 -*-
"""
山洪一张图 空间数据导入脚本（PostgreSQL + PostGIS）

数据源：省平台数据汇总_V20260208_修订版_合并之后直接入库格式（不含隐患tif）
  - 126个小流域/126个小流域.shp        -> sh_map_watershed
  - 空间数据/断面平面位置.geojson       -> sh_map_section
  - 空间数据/防治对象.geojson           -> sh_map_prevention_area
  - 空间数据/隐患要素分布.geojson       -> sh_map_hazard_point
  - 风险隐患/1.山洪灾害防治对象名录.xlsx -> sh_survey_object
  - 风险隐患/2.跨沟道路桥涵塘坝.xlsx     -> sh_survey_road
  - 风险隐患/3.沟滩占地调查成果表.xlsx   -> sh_survey_gully
  - 河道横纵断面成果表/沟道横断面成果表.xlsx -> sh_section_cross_point
  - 河道横纵断面成果表/沟道纵断面成果表.xlsx -> sh_section_long_point

用法：
  python import_shmap.py --dry-run                 # 只解析出统计报告，不连库
  python import_shmap.py --host 127.0.0.1 --db shreport_map --user postgres --password xxx
  python import_shmap.py --rebuild ...             # 导入前清空 9 张表
报告输出：import_report.txt
"""
import argparse
import json
import os
import re
import sys
import time

DEFAULT_DATA_DIR = r"F:\省平台数据上报\省平台数据汇总_V20260208_修订版_合并之后直接入库格式（不含隐患tif）"

import shapefile  # pyshp
import openpyxl

# ---------------------------------------------------------------- 常量

RISK_LABELS = [
    ("risk_narrow_channel", "束窄"),
    ("risk_bend_channel", "急弯"),
    ("risk_depression", "低洼地"),
    ("risk_near_river_landslide", "临河滑坡"),
    ("risk_debris_flow", "泥石流"),
    ("Goutan_oland", "沟滩占地"),
    ("risk_brust", "溃决"),
    ("risk_backwater", "壅水"),
    ("risk_block", "顶托"),
    ("risk_diversion", "改道"),
    ("risk_overflow", "漫流"),
    ("risk_landslide", "潜在滑坡"),
]

# 隐患/防治对象类型归一（先去空白与半角括号，再按此表归并）
TYPE_SYNONYMS = {
    "跨沟桥梁": "跨沟桥涵",
    "跨沟路涵": "跨沟道路",
    "路涵道路": "跨沟道路",
    "临河滑坡体": "临河滑坡",
    "塘(堰)坝": "塘（堰）坝",
}

REPORT_LINES = []


def report(msg=""):
    print(msg)
    REPORT_LINES.append(str(msg))


# ---------------------------------------------------------------- 工具

def to_str(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "null"):
        return None
    return s


def to_float(v):
    s = to_str(v)
    if s is None:
        return None
    try:
        return float(str(s).replace(",", ""))
    except ValueError:
        return None


def to_int(v):
    f = to_float(v)
    return None if f is None else int(round(f))


def norm_type(v):
    """类型清洗：去空白/换行，半角括号转全角，同义词归并。"""
    s = to_str(v)
    if s is None:
        return None
    s = re.sub(r"\s+", "", s).replace("(", "（").replace(")", "）")
    if len(s) % 2 == 0 and len(s) >= 2 and s[: len(s) // 2] * 2 == s:
        s = s[: len(s) // 2]  # 换行重复折叠：企业\n企业 -> 企业
    return TYPE_SYNONYMS.get(s, s)



# 人工核定的编码映射（sh_report_import/check_unmapped.py 空间包含核对生成，
# 全部为 100% 落入基准流域面的要素归属修正，详见 mapping_proposal.txt）
MANUAL_WSCD_MAP = {
    "WDA380012X200001": "WDA380012X200000",
    "WDA380012X200002": "WDA380012X200000",
    "WDA380012X200003": "WDA380012X200000",
    "WDA380012X200004": "WDA380012X200000",
    "WDA380012X200005": "WDA380012X200000",
    "WDA380012X200006": "WDA380012X200000",
    "WDA380012X200007": "WDA380012X200000",
    "WDA380012X200008": "WDA380012X200000",
    "WDA380012X200009": "WDA380012X200000",
    "WDA380012X200010": "WDA380012X200000",
    "WDA380012X200011": "WDA380012X200000",
    "WDA380012X200012": "WDA380012X200000",
    "WDA380012X200013": "WDA380012X200000",
    "WDA380012X200014": "WDA380012X200000",
    "WDA380012X200015": "WDA380012X200000",
    "WDA380012X200016": "WDA380012X200000",
    "WDA380012X200017": "WDA380012X200000",
    "WDA380012X200018": "WDA380012X200000",
    "WDA380012X200019": "WDA380012X200000",
    "WDA380012X200020": "WDA380012X200000",
    "WDA380012X200021": "WDA380012X200000",
    "WDA380012X200022": "WDA380012X200000",
    "WDA380012X200023": "WDA380012X200000",
    "WDA380012X200024": "WDA380012X200000",
    "WDA380012X200025": "WDA380012X200000",
    "WDA380012X200026": "WDA380012X200000",
    "WDA380012X200027": "WDA380012X200000",
    "WDA380012X200028": "WDA380012X200000",
    "WDA380012X200029": "WDA380012X200000",
    "WDA380012X200030": "WDA380012X200000",
    "WFD160000281EB00": "WFD16307G0000000",
    "WFD1630W00000001": "WFD1630W00000000",
    "WFD1630W00000002": "WFD1630W00000000",
    "WFD1630W00000003": "WFD1630W00000000",
    "WFD1630W00000004": "WFD1630W00000000",
    "WFD1630W00000005": "WFD1630W00000000",
    "WFD1630W00000006": "WFD1630W00000000",
    "WFD1630W00000007": "WFD1630W00000000",
    "WFD1630W00000008": "WFD1630W00000000",
    "WFD1630W00000009": "WFD1630W00000000",
    "WFD1630W00000010": "WFD1630W00000000",
    "WFD1630W00000011": "WFD1630W00000000",
    "WFD1630W00000012": "WFD1630W00000000",
    "WFD1630W00000013": "WFD1630W00000000",
    "WFD1630W00000014": "WFD1630W00000000",
    "WFD1630W00000015": "WFD1630W00000000",
    "WFD1630W00000016": "WFD1630W00000000",
    "WFD1630W00000017": "WFD1630W00000000",
    "WFD1630W00000018": "WFD1630W00000000",
    "WFD1630W00000019": "WFD1630W00000000",
    "WFD1630W00000020": "WFD1630W00000000",
    "WFD1630W00000021": "WFD1630W00000000",
    "WFD1630W00000022": "WFD1630W00000000",
    "WFD1630W00000023": "WFD1630W00000000",
    "WFD1630W00000024": "WFD1630W00000000",
    "WFD1630W00000025": "WFD1630W00000000",
    "WFD1630W00000026": "WFD1630W00000000",
    "WFD1630W00000027": "WFD1630W00000000",
    "WFD1630W00000028": "WFD1630W00000000",
    "WFD1630W00000029": "WFD1630W00000000",
    "WFD1630W00000030": "WFD1630W00000000",
    "WFD1630W00000031": "WFD1630W00000000",
    "WFD1630W00000032": "WFD1630W00000000",
    "WFD1630W00000033": "WFD1630W00000000",
    "WFD1630W00000034": "WFD1630W00000000",
    "WFD1630W00000035": "WFD1630W00000000"
}


class WscdCleaner:
    """小流域编码清洗：大写、去空白；GS/SW 前缀剥离后匹配 126 基准编码。"""

    def __init__(self, base_codes):
        self.base = {c.upper() for c in base_codes}
        self.unmapped = {}

    def clean(self, raw):
        s = to_str(raw)
        if s is None:
            return None
        s = s.strip().upper()
        if s in self.base:
            return s
        for prefix in ("GS", "SW"):
            if s.startswith(prefix) and s[len(prefix):] in self.base:
                return s[len(prefix):]
        if s in MANUAL_WSCD_MAP:
            return MANUAL_WSCD_MAP[s]
        self.unmapped[s] = self.unmapped.get(s, 0) + 1
        return s


def shp_to_geojson_geom(shp):
    """pyshp shape -> GeoJSON Polygon/MultiPolygon coordinates（ring 闭合）。"""
    parts = list(shp.parts) + [len(shp.points)]
    rings = []
    for i in range(len(parts) - 1):
        pts = shp.points[parts[i]:parts[i + 1]]
        if len(pts) < 3:
            continue
        ring = [[p[0], p[1]] for p in pts]
        if ring[0] != ring[-1]:
            ring.append(ring[0])
        rings.append(ring)
    if not rings:
        return None
    if len(rings) == 1:
        return {"type": "Polygon", "coordinates": rings}
    return {"type": "MultiPolygon", "coordinates": [[r] for r in rings]}


def geom_bbox(geom):
    """GeoJSON geometry -> 外包矩形 (xmin, ymin, xmax, ymax)，空几何返回 None。"""
    coords = geom.get("coordinates") if geom else None

    def walk(c, box):
        if isinstance(c[0], (int, float)):
            if box is None:
                box = [c[0], c[1], c[0], c[1]]
            else:
                box[0] = min(box[0], c[0]); box[1] = min(box[1], c[1])
                box[2] = max(box[2], c[0]); box[3] = max(box[3], c[1])
        else:
            for sub in c:
                box = walk(sub, box)
        return box

    try:
        return walk(coords, None)
    except (TypeError, IndexError):
        return None


def xlsx_rows(path, sheet=None):
    """read_only 迭代 xlsx 数据行（values_only）。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    for row in ws.iter_rows(values_only=True):
        yield row
    wb.close()


# ---------------------------------------------------------------- 各数据源

def load_watershed(data_dir, cleaner, stats):
    path = os.path.join(data_dir, "126个小流域", "126个小流域.shp")
    r = shapefile.Reader(path, encoding="gbk")
    fields = [f[0] for f in r.fields if f[0] != "DeletionFlag"]
    rows = []
    seen = set()
    for shp, rec in zip(r.shapes(), r.records()):
        d = dict(zip(fields, rec))
        wscd = cleaner.clean(d.get("WSWSCD"))
        geom = shp_to_geojson_geom(shp)
        if wscd is None or geom is None:
            stats["watershed_skipped"] += 1
            continue
        if wscd in seen:
            stats["watershed_dup"] = stats.get("watershed_dup", 0) + 1
            stats.setdefault("watershed_dup_codes", []).append(wscd)
            continue
        seen.add(wscd)
        rows.append({
            "wscd": wscd,
            "wsnm": to_str(d.get("WKMM")) or to_str(d.get("MC")) or wscd,
            "wsarea": to_float(d.get("WSAREA")),
            "wsperi": to_float(d.get("WSPERI")),
            "wsslp": to_float(d.get("WSSLP")),
            "center_lng": to_float(d.get("CENTERX")),
            "center_lat": to_float(d.get("CENTERY")),
            "center_elv": to_float(d.get("CENTERELV")),
            "outlet_lng": to_float(d.get("OUTLETX")),
            "outlet_lat": to_float(d.get("OUTLETY")),
            "outlet_elv": to_float(d.get("OUTLETELV")),
            "max_elv": to_float(d.get("MAXELV")),
            "river_name": to_str(d.get("NAME")),
            "alias_name": to_str(d.get("重命名")),
            "jizhen_count": to_int(d.get("集镇个数")),
            "village_count": to_int(d.get("防治村个数")),
            "gully_count": to_int(d.get("山洪沟道条")),
            "reservoirs": to_str(d.get("涉及水库")),
            "attrs": json.dumps(d, ensure_ascii=False, default=str),
            "geom_json": json.dumps(geom, ensure_ascii=False),
        })
    stats["watershed"] = len(rows)
    return rows


def load_geojson_layer(path, stats_key, stats, wscd_field="wswscd"):
    with open(path, encoding="utf-8") as f:
        g = json.load(f)
    rows = []
    for feat in g.get("features", []):
        props = feat.get("properties") or {}
        geom = feat.get("geometry")
        if not geom or geom.get("coordinates") is None:
            stats[stats_key + "_skipped"] += 1
            continue
        row = dict(props)
        row["wscd_raw"] = props.get(wscd_field)
        row["geom_json"] = json.dumps(geom, ensure_ascii=False)
        row["geom_type"] = geom.get("type")
        rows.append(row)
    stats[stats_key] = len(rows)
    return rows


def load_named_sheet(path, cols, stats_key, stats):
    """风险隐患 xlsx：第1行中文表头，第2行英文字段名，第3行起为数据。
    cols: {列下标: 业务字段名}"""
    rows = []
    it = xlsx_rows(path)
    next(it)  # 中文表头
    next(it)  # 英文字段名
    for i, row in enumerate(it, start=1):
        d = {key: row[idx] if idx < len(row) else None for idx, key in cols.items()}
        d["src_id"] = to_int(d.get("src_id"))
        if not any(to_str(v) for v in d.values()):
            continue
        d["_row"] = i
        rows.append(d)
    stats[stats_key] = len(rows)
    return rows


def load_survey_object(data_dir, stats):
    path = os.path.join(data_dir, "风险隐患", "1.山洪灾害防治对象名录_20260209_095441_.xlsx")
    cols = {0: "src_id", 1: "county_name", 2: "county_adcd", 3: "town_name", 4: "town_adcd",
            5: "fnm", 6: "fcd", 7: "ftp", 8: "pcount", 9: "rvnm", 10: "rvcd",
            11: "is_road_cro", 12: "is_road_code", 13: "is_dam", 14: "is_dam_code",
            15: "is_confluence", 16: "is_confluence_code",
            17: "risk_narrow_channel", 18: "risk_bend_channel", 19: "risk_depression",
            20: "risk_near_river_landslide", 21: "risk_debris_flow", 22: "Goutan_oland",
            23: "risk_brust", 24: "risk_backwater", 25: "risk_block", 26: "risk_diversion",
            27: "risk_overflow", 28: "comments", 29: "wsnm", 30: "wscd", 31: "risk_landslide",
            32: "status", 33: "survey_year", 34: "reporter", 35: "report_time",
            36: "base_lgtd", 37: "base_lttd", 38: "elevation", 39: "river_section_name",
            40: "lgtd", 41: "lttd", 42: "type"}
    raw = load_named_sheet(path, cols, "survey_object", stats)
    risk_keys = [k for k, _ in RISK_LABELS]
    rows = []
    for d in raw:
        flags = []
        for k, label in RISK_LABELS:
            v = to_str(d.get(k))
            if v and v not in ("0", "否", "nan"):
                flags.append(label)
        rows.append({
            "src_id": d["src_id"],
            "county_name": to_str(d.get("county_name")), "county_adcd": to_str(d.get("county_adcd")),
            "town_name": to_str(d.get("town_name")), "town_adcd": to_str(d.get("town_adcd")),
            "fnm": to_str(d.get("fnm")), "fcd": to_str(d.get("fcd")), "ftp": norm_type(d.get("ftp")),
            "pcount": to_int(d.get("pcount")),
            "rvnm": to_str(d.get("rvnm")), "rvcd": to_str(d.get("rvcd")),
            "wsnm": to_str(d.get("wsnm")), "wscd": to_str(d.get("wscd")),
            "lgtd": to_float(d.get("lgtd")), "lttd": to_float(d.get("lttd")),
            "base_lgtd": to_float(d.get("base_lgtd")), "base_lttd": to_float(d.get("base_lttd")),
            "elevation": to_float(d.get("elevation")),
            "river_section_name": to_str(d.get("river_section_name")),
            "comments": to_str(d.get("comments")),
            "risk_flags": flags,
            "attrs": json.dumps(d, ensure_ascii=False, default=str),
        })
    stats["survey_object_flags_avg"] = (sum(len(r["risk_flags"]) for r in rows) / len(rows)) if rows else 0
    return rows


def load_survey_road(data_dir, stats):
    path = os.path.join(data_dir, "风险隐患", "2.跨沟道路、桥涵、塘坝调查成果表_20260209_095446_.xlsx")
    cols = {0: "src_id", 1: "county_name", 2: "county_adcd", 3: "town_name", 4: "town_adcd",
            5: "name", 6: "code", 7: "lgtd", 8: "lttd", 9: "stype", 10: "channelw", 11: "channeld",
            12: "csshpt", 13: "obstruction_area_ratio", 14: "obstruction_storage_capacity",
            15: "rvnm", 16: "rv_code", 17: "backwater_objname", 18: "backwater_objcode",
            19: "burst_objname", 20: "burst_objcode", 21: "overflow_objname", 22: "overflow_objcode",
            23: "comments", 24: "wsnm", 25: "wscd", 26: "adnm", 27: "adcd", 28: "big_category",
            29: "sub_category", 30: "rvnm2", 31: "dam_len", 32: "dam_wid", 33: "length",
            34: "width", 35: "height", 36: "status", 37: "survery_year", 38: "creator", 39: "creator_at"}
    raw = load_named_sheet(path, cols, "survey_road", stats)
    rows = []
    for d in raw:
        rows.append({
            "src_id": d["src_id"],
            "county_name": to_str(d.get("county_name")), "county_adcd": to_str(d.get("county_adcd")),
            "town_name": to_str(d.get("town_name")), "town_adcd": to_str(d.get("town_adcd")),
            "name": to_str(d.get("name")), "code": to_str(d.get("code")),
            "lgtd": to_float(d.get("lgtd")), "lttd": to_float(d.get("lttd")),
            "stype": to_str(d.get("stype")),
            "channelw": to_str(d.get("channelw")), "channeld": to_str(d.get("channeld")),
            "csshpt": to_str(d.get("csshpt")),
            "obstruction_area_ratio": to_str(d.get("obstruction_area_ratio")),
            "obstruction_storage_capacity": to_float(d.get("obstruction_storage_capacity")),
            "rvnm": to_str(d.get("rvnm")) or to_str(d.get("rvnm2")), "rv_code": to_str(d.get("rv_code")),
            "backwater_objname": to_str(d.get("backwater_objname")), "backwater_objcode": to_str(d.get("backwater_objcode")),
            "burst_objname": to_str(d.get("burst_objname")), "burst_objcode": to_str(d.get("burst_objcode")),
            "overflow_objname": to_str(d.get("overflow_objname")), "overflow_objcode": to_str(d.get("overflow_objcode")),
            "comments": to_str(d.get("comments")),
            "wsnm": to_str(d.get("wsnm")), "wscd": to_str(d.get("wscd")),
            "attrs": json.dumps(d, ensure_ascii=False, default=str),
        })
    return rows


def load_survey_gully(data_dir, stats):
    path = os.path.join(data_dir, "风险隐患", "3.沟滩占地情况调查成果表_20260209_095449_.xlsx")
    cols = {0: "src_id", 1: "county_name", 2: "county_adcd", 3: "town_name", 4: "town_adcd",
            5: "name", 6: "code", 7: "lgtd", 8: "lttd", 9: "stype", 10: "width", 11: "deep",
            12: "dmform", 13: "zsmj", 14: "rvnm", 15: "rvcd", 16: "comments", 17: "wsnm", 18: "wscd"}
    raw = load_named_sheet(path, cols, "survey_gully", stats)
    rows = []
    for d in raw:
        rows.append({
            "src_id": d["src_id"],
            "county_name": to_str(d.get("county_name")), "county_adcd": to_str(d.get("county_adcd")),
            "town_name": to_str(d.get("town_name")), "town_adcd": to_str(d.get("town_adcd")),
            "name": to_str(d.get("name")), "code": to_str(d.get("code")),
            "lgtd": to_float(d.get("lgtd")), "lttd": to_float(d.get("lttd")),
            "stype": to_str(d.get("stype")),
            "width": to_str(d.get("width")), "deep": to_str(d.get("deep")),
            "dmform": to_str(d.get("dmform")), "zsmj": to_str(d.get("zsmj")),
            "rvnm": to_str(d.get("rvnm")), "rvcd": to_str(d.get("rvcd")),
            "comments": to_str(d.get("comments")),
            "wsnm": to_str(d.get("wsnm")), "wscd": to_str(d.get("wscd")),
            "attrs": json.dumps(d, ensure_ascii=False, default=str),
        })
    return rows


def load_section_points(path, kind, stats):
    """断面成果表：仅第1行中文表头，第2行起为数据。kind: cross|long"""
    cols = ({"cross": {0: "point_no", 1: "point_desc", 2: "dist", 3: "elev", 4: "lgtd", 5: "lttd",
                       6: "roughness", 7: "adcd", 8: "rvcd", 9: "pid", 10: "wscd", 11: "wsnm"},
             "long": {0: "point_no", 1: "point_desc", 2: "dist", 3: "dist_dir", 4: "bed_elev",
                      5: "water_elev", 6: "lgtd", 7: "lttd", 8: "rvcd", 9: "adcd", 10: "pid",
                      11: "wscd", 12: "wsnm"}})[kind]
    rows = []
    pids = set()
    it = xlsx_rows(path)
    next(it)  # 表头
    for row in it:
        d = {key: row[idx] if idx < len(row) else None for idx, key in cols.items()}
        pid = to_str(d.get("pid"))
        if pid is None:
            continue
        pids.add(pid)
        item = {"pid": pid, "point_no": to_int(d.get("point_no")),
                "point_desc": to_str(d.get("point_desc")),
                "dist": to_float(d.get("dist")), "lgtd": to_float(d.get("lgtd")),
                "lttd": to_float(d.get("lttd")), "rvcd": to_str(d.get("rvcd")),
                "adcd": to_str(d.get("adcd")), "wscd": to_str(d.get("wscd")),
                "wsnm": to_str(d.get("wsnm"))}
        if kind == "cross":
            item["elev"] = to_float(d.get("elev"))
            item["roughness"] = to_float(d.get("roughness"))
        else:
            item["dist_dir"] = to_float(d.get("dist_dir"))
            item["bed_elev"] = to_float(d.get("bed_elev"))
            item["water_elev"] = to_float(d.get("water_elev"))
        rows.append(item)
    stats[f"section_{kind}"] = len(rows)
    stats[f"section_{kind}_pids"] = len(pids)
    return rows, pids


# ---------------------------------------------------------------- 入库

INSERTS = {
    "sh_map_watershed": """INSERT INTO sh_map_watershed (wscd, wsnm, wsarea, wsperi, wsslp,
        center_lng, center_lat, center_elv, outlet_lng, outlet_lat, outlet_elv, max_elv,
        river_name, alias_name, jizhen_count, village_count, gully_count, reservoirs, attrs, geom)
        SELECT wscd, wsnm, wsarea, wsperi, wsslp, center_lng, center_lat, center_elv,
        outlet_lng, outlet_lat, outlet_elv, max_elv, river_name, alias_name,
        jizhen_count, village_count, gully_count, reservoirs, attrs::jsonb,
        ST_Multi(ST_SetSRID(ST_GeomFromGeoJSON(geom), 4326))
        FROM (VALUES %s) AS v(wscd, wsnm, wsarea, wsperi, wsslp, center_lng, center_lat,
        center_elv, outlet_lng, outlet_lat, outlet_elv, max_elv, river_name, alias_name,
        jizhen_count, village_count, gully_count, reservoirs, attrs, geom)""",
    "sh_map_section": """INSERT INTO sh_map_section (pid, sec_type, sec_name, rvname, rvcd, wscd, geom)
        SELECT pid, sec_type, sec_name, rvname, rvcd, wscd,
        ST_Multi(ST_SetSRID(ST_GeomFromGeoJSON(geom), 4326))
        FROM (VALUES %s) AS v(pid, sec_type, sec_name, rvname, rvcd, wscd, geom)""",
    "sh_map_prevention_area": """INSERT INTO sh_map_prevention_area (obj_pid, obj_name, obj_type, pcount, rvname, rvcd, wscd, geom)
        SELECT obj_pid, obj_name, obj_type, pcount, rvname, rvcd, wscd,
        ST_Multi(ST_SetSRID(ST_GeomFromGeoJSON(geom), 4326))
        FROM (VALUES %s) AS v(obj_pid, obj_name, obj_type, pcount, rvname, rvcd, wscd, geom)""",
    "sh_map_hazard_point": """INSERT INTO sh_map_hazard_point (pid, name, hz_type, rvname, rvcd, wscd, fcd, fnm, lgtd, lttd, geom)
        SELECT pid, name, hz_type, rvname, rvcd, wscd, fcd, fnm, lgtd, lttd,
        ST_SetSRID(ST_GeomFromGeoJSON(geom), 4326)
        FROM (VALUES %s) AS v(pid, name, hz_type, rvname, rvcd, wscd, fcd, fnm, lgtd, lttd, geom)""",
    "sh_survey_object": """INSERT INTO sh_survey_object (src_id, county_name, county_adcd, town_name, town_adcd,
        fnm, fcd, ftp, pcount, rvnm, rvcd, wsnm, wscd, lgtd, lttd, base_lgtd, base_lttd, elevation,
        river_section_name, comments, risk_flags, attrs) VALUES %s""",
    "sh_survey_road": """INSERT INTO sh_survey_road (src_id, county_name, county_adcd, town_name, town_adcd,
        name, code, lgtd, lttd, stype, channelw, channeld, csshpt, obstruction_area_ratio,
        obstruction_storage_capacity, rvnm, rv_code, backwater_objname, backwater_objcode,
        burst_objname, burst_objcode, overflow_objname, overflow_objcode, comments, wsnm, wscd, attrs)
        VALUES %s""",
    "sh_survey_gully": """INSERT INTO sh_survey_gully (src_id, county_name, county_adcd, town_name, town_adcd,
        name, code, lgtd, lttd, stype, width, deep, dmform, zsmj, rvnm, rvcd, comments, wsnm, wscd, attrs)
        VALUES %s""",
    "sh_section_cross_point": """INSERT INTO sh_section_cross_point (pid, point_no, point_desc, dist, elev, lgtd, lttd,
        roughness, rvcd, adcd, wscd, wsnm) VALUES %s""",
    "sh_section_long_point": """INSERT INTO sh_section_long_point (pid, point_no, point_desc, dist, dist_dir, bed_elev,
        water_elev, lgtd, lttd, rvcd, adcd, wscd, wsnm) VALUES %s""",
}

ALL_TABLES = list(INSERTS.keys())


def pg_insert(cur, table, values_seq, page_size=2000):
    from psycopg2.extras import execute_values
    total = 0
    for i in range(0, len(values_seq), page_size):
        chunk = values_seq[i:i + page_size]
        execute_values(cur, INSERTS[table], chunk, page_size=len(chunk))
        total += len(chunk)
    return total


def run_import(args, data, stats):
    import psycopg2
    conn = psycopg2.connect(host=args.host, port=args.port, dbname=args.db,
                            user=args.user, password=args.password)
    conn.autocommit = False
    cur = conn.cursor()
    ddl = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "pgsql_ddl.sql"),
               encoding="utf-8").read()
    if args.rebuild:
        cur.execute("DROP TABLE IF EXISTS " + ", ".join(t for t in ALL_TABLES) + " CASCADE")
        conn.commit()
        report("[DB] --rebuild：9 张旧表已删除，将按最新 DDL 重建")

    cur.execute(ddl)
    conn.commit()
    report("[DB] DDL 执行完成")

    t0 = time.time()
    counts = {}

    def bbox_of(geom_json):
        return geom_bbox(json.loads(geom_json)) if geom_json else None

    # 1 流域面
    ws_vals = []
    for r in data["watershed"]:
        bb = bbox_of(r["geom_json"])
        ws_vals.append((
            r["wscd"], r["wsnm"], r["wsarea"], r["wsperi"], r["wsslp"],
            r["center_lng"], r["center_lat"], r["center_elv"],
            r["outlet_lng"], r["outlet_lat"], r["outlet_elv"], r["max_elv"],
            r["river_name"], r["alias_name"], r["jizhen_count"], r["village_count"],
            r["gully_count"], r["reservoirs"], r["attrs"], r["geom_json"],
            bb[0], bb[1], bb[2], bb[3]))
    counts["sh_map_watershed"] = pg_insert(cur, "sh_map_watershed", ws_vals)
    conn.commit()
    # 2 断面线
    sec_vals = []
    for r in data["sections"]:
        bb = bbox_of(r["geom_json"])
        sec_vals.append((
            to_str(r.get("pid")), to_str(r.get("types")), to_str(r.get("name")),
            to_str(r.get("rvname")), to_str(r.get("rvcd")),
            data["cleaner"].clean(r.get("wscd_raw")), r["geom_json"],
            bb[0], bb[1], bb[2], bb[3]))
    counts["sh_map_section"] = pg_insert(cur, "sh_map_section", sec_vals)
    conn.commit()
    # 3 防治对象面（pid 换行重复取第一个，rvcd 逗号拼接取第一个）
    po_vals = []
    for r in data["preventions"]:
        bb = bbox_of(r["geom_json"])
        pid = to_str(r.get("pid"))
        if pid and "\n" in pid:
            pid = re.split(r"[\r\n]+", pid)[0]
        rvcd = to_str(r.get("rvcd"))
        if rvcd and "," in rvcd:
            rvcd = rvcd.split(",")[0]
        po_vals.append((
            pid, to_str(r.get("name")), norm_type(r.get("types")),
            to_int(r.get("pcount")), to_str(r.get("rvname")), rvcd,
            data["cleaner"].clean(r.get("wscd_raw")), r["geom_json"],
            bb[0], bb[1], bb[2], bb[3]))
    counts["sh_map_prevention_area"] = pg_insert(cur, "sh_map_prevention_area", po_vals)
    conn.commit()
    # 4 隐患点（属性列 lgtd/lttd 多为空，从 Point 几何坐标提取；源数据存在
    # 经纬度颠倒/重复的脏点，超出合理范围时置空，避免撑爆地图视野）
    hz_vals = []
    for r in data["hazards"]:
        geom = json.loads(r["geom_json"])
        coords = geom.get("coordinates") if geom.get("type") == "Point" else None
        lgtd, lttd = (coords[0], coords[1]) if coords and len(coords) == 2 else (to_float(r.get("lgtd")), to_float(r.get("lttd")))
        if lgtd is None or lttd is None or lgtd == lttd or not (89 <= lgtd <= 112) or not (30 <= lttd <= 45):
            lgtd = lttd = None
        hz_vals.append((
            to_str(r.get("pid")), to_str(r.get("name")), norm_type(r.get("types")),
            to_str(r.get("rvname")), to_str(r.get("rvcd")),
            data["cleaner"].clean(r.get("wscd_raw")),
            to_str(r.get("fcd")), to_str(r.get("fnm")),
            lgtd, lttd, r["geom_json"]))
    counts["sh_map_hazard_point"] = pg_insert(cur, "sh_map_hazard_point", hz_vals)
    conn.commit()
    # 5-7 调查表
    counts["sh_survey_object"] = pg_insert(cur, "sh_survey_object", [(
        r["src_id"], r["county_name"], r["county_adcd"], r["town_name"], r["town_adcd"],
        r["fnm"], r["fcd"], r["ftp"], r["pcount"], r["rvnm"], r["rvcd"], r["wsnm"], r["wscd"],
        r["lgtd"], r["lttd"], r["base_lgtd"], r["base_lttd"], r["elevation"],
        r["river_section_name"], r["comments"], r["risk_flags"], r["attrs"]) for r in data["survey_object"]])
    conn.commit()
    counts["sh_survey_road"] = pg_insert(cur, "sh_survey_road", [(
        r["src_id"], r["county_name"], r["county_adcd"], r["town_name"], r["town_adcd"],
        r["name"], r["code"], r["lgtd"], r["lttd"], r["stype"], r["channelw"], r["channeld"],
        r["csshpt"], r["obstruction_area_ratio"], r["obstruction_storage_capacity"],
        r["rvnm"], r["rv_code"], r["backwater_objname"], r["backwater_objcode"],
        r["burst_objname"], r["burst_objcode"], r["overflow_objname"], r["overflow_objcode"],
        r["comments"], r["wsnm"], r["wscd"], r["attrs"]) for r in data["survey_road"]])
    conn.commit()
    counts["sh_survey_gully"] = pg_insert(cur, "sh_survey_gully", [(
        r["src_id"], r["county_name"], r["county_adcd"], r["town_name"], r["town_adcd"],
        r["name"], r["code"], r["lgtd"], r["lttd"], r["stype"], r["width"], r["deep"],
        r["dmform"], r["zsmj"], r["rvnm"], r["rvcd"], r["comments"], r["wsnm"], r["wscd"],
        r["attrs"]) for r in data["survey_gully"]])
    conn.commit()
    # 8-9 断面成果点
    counts["sh_section_cross_point"] = pg_insert(cur, "sh_section_cross_point", [(
        r["pid"], r["point_no"], r["point_desc"], r["dist"], r["elev"], r["lgtd"], r["lttd"],
        r["roughness"], r["rvcd"], r["adcd"], r["wscd"], r["wsnm"]) for r in data["cross"]])
    conn.commit()
    counts["sh_section_long_point"] = pg_insert(cur, "sh_section_long_point", [(
        r["pid"], r["point_no"], r["point_desc"], r["dist"], r["dist_dir"], r["bed_elev"],
        r["water_elev"], r["lgtd"], r["lttd"], r["rvcd"], r["adcd"], r["wscd"], r["wsnm"]) for r in data["long"]])
    conn.commit()

    report(f"[DB] 入库完成，耗时 {time.time() - t0:.1f}s")
    for t, c in counts.items():
        report(f"  {t}: {c} 行")
    cur.close()
    conn.close()


# ---------------------------------------------------------------- 主流程

def main():
    ap = argparse.ArgumentParser(description="山洪一张图空间数据导入 PostgreSQL+PostGIS")
    ap.add_argument("--data-dir", default=DEFAULT_DATA_DIR)
    ap.add_argument("--host", default=os.environ.get("PGHOST", "127.0.0.1"))
    ap.add_argument("--port", default=os.environ.get("PGPORT", "5432"))
    ap.add_argument("--db", default=os.environ.get("PGDATABASE", "shreport_map"))
    ap.add_argument("--user", default=os.environ.get("PGUSER", "postgres"))
    ap.add_argument("--password", default=os.environ.get("PGPASSWORD", ""))
    ap.add_argument("--dry-run", action="store_true", help="只解析出统计报告，不连库")
    ap.add_argument("--rebuild", action="store_true", help="导入前清空 9 张表")
    args = ap.parse_args()

    data_dir = args.data_dir
    stats = {"watershed_skipped": 0, "section_skipped": 0,
             "prevention_area_skipped": 0, "hazard_point_skipped": 0}
    report("=" * 72)
    report("山洪一张图空间数据导入 " + ("【干跑模式，不连库】" if args.dry_run else "【入库模式】"))
    report("数据目录: " + data_dir)
    report("=" * 72)

    t0 = time.time()
    # 先读流域面拿到基准编码集
    pre_reader = shapefile.Reader(os.path.join(data_dir, "126个小流域", "126个小流域.shp"), encoding="gbk")
    wscd_field_idx = [f[0] for f in pre_reader.fields if f[0] != "DeletionFlag"].index("WSWSCD")
    base_codes = [rec[wscd_field_idx] for rec in pre_reader.records()]
    cleaner = WscdCleaner(base_codes)
    stats["base_codes"] = len(base_codes)

    data = {}
    data["watershed"] = load_watershed(data_dir, cleaner, stats)
    data["sections"] = load_geojson_layer(os.path.join(data_dir, "空间数据", "断面平面位置.geojson"),
                                          "section", stats)
    data["preventions"] = load_geojson_layer(os.path.join(data_dir, "空间数据", "防治对象.geojson"),
                                             "prevention_area", stats)
    data["hazards"] = load_geojson_layer(os.path.join(data_dir, "空间数据", "隐患要素分布.geojson"),
                                         "hazard_point", stats)
    data["survey_object"] = load_survey_object(data_dir, stats)
    data["survey_road"] = load_survey_road(data_dir, stats)
    data["survey_gully"] = load_survey_gully(data_dir, stats)
    data["cross"], cross_pids = load_section_points(
        os.path.join(data_dir, "河道横纵断面成果表", "沟道横断面成果表.xlsx"), "cross", stats)
    data["long"], long_pids = load_section_points(
        os.path.join(data_dir, "河道横纵断面成果表", "沟道纵断面成果表.xlsx"), "long", stats)
    data["cleaner"] = cleaner

    # ---------------- 统计报告
    section_pids = {to_str(r.get("pid")) for r in data["sections"]}
    report("")
    report("【解析统计】")
    report(f"  基准小流域编码（126个小流域.shp）: {stats['base_codes']}")
    report(f"  sh_map_watershed        流域面: {stats['watershed']}（跳过 {stats['watershed_skipped']}，重复编码去重 {stats.get('watershed_dup', 0)}）")
    if stats.get("watershed_dup_codes"):
        report(f"    重复编码: {sorted(set(stats['watershed_dup_codes']))}")
    report(f"  sh_map_section          断面线: {stats['section']}（跳过 {stats['section_skipped']}）")
    by_type = {}
    for r in data["sections"]:
        by_type[to_str(r.get("types")) or "未知"] = by_type.get(to_str(r.get("types")) or "未知", 0) + 1
    report(f"    断面类型分布: {by_type}")
    report(f"  sh_map_prevention_area  防治对象面: {stats['prevention_area']}（跳过 {stats['prevention_area_skipped']}）")
    by_type = {}
    for r in data["preventions"]:
        t = norm_type(r.get("types")) or "未知"
        by_type[t] = by_type.get(t, 0) + 1
    report(f"    防治类型分布（清洗后）: {by_type}")
    report(f"  sh_map_hazard_point     隐患点: {stats['hazard_point']}（跳过 {stats['hazard_point_skipped']}）")
    by_type = {}
    for r in data["hazards"]:
        t = norm_type(r.get("types")) or "未知"
        by_type[t] = by_type.get(t, 0) + 1
    report(f"    隐患类型分布（清洗后）: {by_type}")
    report(f"  sh_survey_object        防治对象名录: {stats['survey_object']}（人均风险标志 {stats.get('survey_object_flags_avg', 0):.2f} 个）")
    report(f"  sh_survey_road          跨沟道路桥涵塘坝: {stats['survey_road']}")
    report(f"  sh_survey_gully         沟滩占地: {stats['survey_gully']}")
    report(f"  sh_section_cross_point  横断面成果点: {stats['section_cross']}（断面 {stats['section_cross_pids']} 条）")
    report(f"  sh_section_long_point   纵断面成果点: {stats['section_long']}（断面 {stats['section_long_pids']} 条）")

    report("")
    report("【断面编码关联】（断面平面位置.pid ↔ 成果表.断面编码）")
    report(f"  断面线 pid 有横断面成果: {len(section_pids & cross_pids)} / {len(section_pids)}")
    report(f"  断面线 pid 有纵断面成果: {len(section_pids & long_pids)} / {len(section_pids)}")
    only_cross = cross_pids - section_pids
    only_long = long_pids - section_pids
    report(f"  仅有成果无线要素的断面: 横 {len(only_cross)} 条, 纵 {len(only_long)} 条")
    if only_cross:
        report(f"    横断面示例: {sorted(only_cross)[:5]}")
    if only_long:
        report(f"    纵断面示例: {sorted(only_long)[:5]}")

    report("")
    report("【小流域编码清洗】")
    for name, rows in (("断面平面位置", data["sections"]), ("防治对象", data["preventions"]),
                       ("隐患要素分布", data["hazards"])):
        mapped, raws = 0, set()
        for r in rows:
            c = cleaner.clean(r.get("wscd_raw"))
            if c:
                mapped += 1
            raws.add(to_str(r.get("wscd_raw")))
        in_base = sum(1 for x in raws if x and x.strip().upper() in cleaner.base)
        report(f"  {name}: 原始编码 {len(raws)} 种 / {mapped} 条，直接命中基准 {in_base} 种")

    if cleaner.unmapped:
        report(f"  未命中 126 基准编码的编码 {len(cleaner.unmapped)} 种（已按原样入库）:")
        for k, v in sorted(cleaner.unmapped.items()):
            report(f"    {k}  x{v}")

    report("")
    report(f"【耗时】解析 {time.time() - t0:.1f}s")

    if args.dry_run:
        report("")
        report("干跑结束：未连接数据库。去掉 --dry-run 并提供连接参数后执行入库。")
    else:
        run_import(args, data, stats)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "import_report.txt")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(REPORT_LINES))
    print("\n报告已写入:", out)


if __name__ == "__main__":
    sys.exit(main())
